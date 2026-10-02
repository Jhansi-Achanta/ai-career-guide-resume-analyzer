"""Tests for the Flask JSON API and the resume parser.

The Gemini calls are mocked, so these tests run without an API key and without
using any network quota.  Real JSON files are used, but they are written to a
temporary folder so your own history in ``data/`` is never touched.

Run them with::

    python -m unittest tests.test_api -v
"""

from __future__ import annotations

import json
import sys
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest import mock

# Make sure "app", "config", "services" and "storage" can be imported.
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import config  # noqa: E402
from app import app  # noqa: E402
from google.genai import errors as genai_errors  # noqa: E402
from services.gemini_client import (  # noqa: E402
    AIRequestError,
    _code_of,
    _run_with_retry,
    generate_json,
)
from services.resume_parser import ResumeParseError, extract_text  # noqa: E402

FAKE_ANALYSIS = {
    "ats_score": 72,
    "score_label": "Good",
    "resume_summary": "A final year student with Python and Flask experience.",
    "score_breakdown": [
        {"category": "Keywords", "score": 20, "max_score": 30, "comment": "Some gaps."}
    ],
    "detected_skills": ["Python", "Flask"],
    "missing_skills": [
        {"skill": "Docker", "why_it_matters": "Used for deployment.", "priority": "High"}
    ],
    "missing_keywords": ["REST API"],
    "certification_suggestions": [
        {
            "name": "AWS Cloud Practitioner",
            "provider": "AWS",
            "why": "Cloud basics.",
            "priority": "Medium",
        }
    ],
    "project_improvement_suggestions": [
        {"title": "Chat app", "what_to_add": "Add tests", "impact": "Shows quality"}
    ],
    "formatting_suggestions": ["Keep to one page."],
    "content_suggestions": ["Add numbers to bullet points."],
    "target_role_comparison": {
        "target_role": "Backend Developer",
        "match_percentage": 65,
        "matching_strengths": ["Python"],
        "gaps": ["Docker"],
        "verdict": "Close, but deployment experience is missing.",
    },
    "job_portal_suggestions": [
        {"name": "Naukri", "url": "https://www.naukri.com", "why": "Large Indian job board."}
    ],
    "action_steps": ["Add a deployment project."],
}

FAKE_PLAN = {
    "career_summary": "You already have a good base for backend work.",
    "recommended_roles": [
        {
            "title": "Backend Developer",
            "why_it_fits": "Python skills",
            "entry_level": "Fresher",
            "growth_path": "Senior Backend",
        }
    ],
    "recommended_skills": [
        {
            "name": "Docker",
            "category": "Tool",
            "why": "Deployment",
            "priority": "High",
            "learn_first": True,
        }
    ],
    "learning_priorities": [
        {"order": 1, "focus": "Python", "reason": "Core language", "estimated_hours": 20}
    ],
    "thirty_day_plan": [
        {
            "period": "Days 1-5",
            "focus": "Python basics",
            "tasks": ["Practise daily"],
            "deliverable": "5 small scripts",
            "estimated_time": "1 hour/day",
        }
    ],
    "learning_resources": [
        {
            "title": "Python Docs",
            "type": "Docs",
            "provider": "Python",
            "url": "https://docs.python.org",
            "why": "Reference",
        }
    ],
    "project_suggestions": [
        {
            "title": "Task API",
            "description": "A REST API",
            "skills_practiced": ["Flask"],
            "difficulty": "Beginner",
        }
    ],
    "preparation_tips": ["Build a portfolio."],
    "interview_preparation": ["Practise DSA."],
    "next_steps": ["Start the 30-day plan."],
}


def make_docx_bytes() -> bytes:
    """Build a small but realistic DOCX resume in memory."""
    import docx

    document = docx.Document()
    paragraphs = [
        "Ravi Kumar",
        "Backend Developer | ravi.kumar@example.com | +91 90000 00000",
        "SUMMARY",
        "Final year B.Tech Computer Science student with hands-on experience building",
        "REST APIs using Python and Flask. Comfortable with MySQL and Git.",
        "SKILLS",
        "Python, Flask, MySQL, HTML, CSS, JavaScript, Git, Postman, problem solving",
        "PROJECTS",
        "Student Result Portal - built a Flask web app with MySQL authentication.",
    ]
    for line in paragraphs:
        document.add_paragraph(line)

    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Education"
    table.cell(0, 1).text = "B.Tech CSE, Anna University, 2027"
    table.cell(1, 0).text = "Certification"
    table.cell(1, 1).text = "Python for Everybody"

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class ApiTestCase(unittest.TestCase):
    """Base class: temporary JSON stores + a Flask test client."""

    def setUp(self) -> None:
        self._tempdir = tempfile.TemporaryDirectory()
        base = Path(self._tempdir.name)

        # Redirect the JSON stores into the temp folder.
        self._original = (config.ANALYSES_FILE, config.CAREER_PLANS_FILE)
        config.ANALYSES_FILE = base / "analyses.json"
        config.CAREER_PLANS_FILE = base / "career_plans.json"

        from storage import records

        records.init_stores()

        app.config["TESTING"] = True
        self.client = app.test_client()

    def tearDown(self) -> None:
        config.ANALYSES_FILE, config.CAREER_PLANS_FILE = self._original
        self._tempdir.cleanup()


class GeminiClientTests(unittest.TestCase):
    def test_rate_limits_are_not_retried_for_any_sdk_error_path(self):
        generic_error = Exception("rate limited")
        generic_error.status_code = 429
        rate_limit_errors = (
            genai_errors.ClientError(code=429, response_json={}),
            genai_errors.APIError(code=429, response_json={}),
            generic_error,
        )

        for error in rate_limit_errors:
            with self.subTest(error_type=type(error).__name__):
                operation = mock.Mock(side_effect=error)
                with mock.patch(
                    "services.gemini_client._RETRY_DELAYS", (0.0, 0.0, 0.0)
                ):
                    with self.assertRaises(AIRequestError) as raised:
                        _run_with_retry(operation)

                self.assertEqual(raised.exception.status_code, 429)
                operation.assert_called_once_with()

    def test_generate_json_uses_interactions_and_parses_output_text(self):
        schema = {
            "type": "object",
            "properties": {"name": {"type": "string"}},
            "required": ["name"],
        }
        fake_client = mock.Mock()
        fake_client.interactions.create.return_value.output_text = '{"name":"Ada"}'

        with mock.patch("services.gemini_client.get_client", return_value=fake_client):
            result = generate_json("Return a name", schema, "JSON only", temperature=0.2)

        self.assertEqual(result, {"name": "Ada"})
        kwargs = fake_client.interactions.create.call_args.kwargs
        self.assertEqual(kwargs["model"], config.GEMINI_MODEL)
        self.assertFalse(kwargs["store"])
        self.assertEqual(kwargs["timeout"], 45)
        self.assertIn('"required": ["name"]', kwargs["input"])
        self.assertNotIn("response_schema", kwargs)

    def test_generate_json_rejects_missing_nested_required_fields(self):
        schema = {
            "type": "object",
            "properties": {
                "items": {
                    "type": "array",
                    "items": {
                        "type": "object",
                        "properties": {"name": {"type": "string"}},
                        "required": ["name"],
                    },
                }
            },
            "required": ["items"],
        }
        fake_client = mock.Mock()
        fake_client.interactions.create.return_value.output_text = '{"items":[{}]}'

        with mock.patch("services.gemini_client.get_client", return_value=fake_client):
            with self.assertRaises(AIRequestError):
                generate_json("Return items", schema, "JSON only")

    def test_interactions_503_becomes_a_clean_api_error(self):
        sdk_error_type = type(
            "InteractionsServerError",
            (Exception,),
            {"__module__": "google.genai._gaos.errors"},
        )
        sdk_error = sdk_error_type("high demand")
        sdk_error.status_code = 503
        self.assertEqual(_code_of(sdk_error), 503)
        fake_client = mock.Mock()
        fake_client.interactions.create.side_effect = sdk_error

        with mock.patch("services.gemini_client.get_client", return_value=fake_client):
            with mock.patch("services.gemini_client._RETRY_DELAYS", (0.0,)):
                with self.assertRaises(AIRequestError) as raised:
                    generate_json("Return a name", {}, "JSON only")

        self.assertEqual(raised.exception.status_code, 503)
        self.assertNotIn("high demand", raised.exception.message)


class HealthTests(ApiTestCase):
    def test_health_never_leaks_the_api_key(self):
        response = self.client.get("/api/health")
        self.assertEqual(response.status_code, 200)
        body = response.get_json()
        self.assertEqual(body["status"], "ok")
        self.assertIn("ai_configured", body)
        self.assertNotIn("api_key", json.dumps(body).lower())

    def test_target_roles(self):
        body = self.client.get("/api/target-roles").get_json()
        self.assertIn("Backend Developer", body["roles"])
        self.assertTrue(body["experience_levels"])


class PageTests(ApiTestCase):
    def test_pages_are_served(self):
        for path in ("/", "/career", "/analyzer", "/history"):
            with self.subTest(path=path):
                response = self.client.get(path)
                self.assertEqual(response.status_code, 200)
                response.close()  # release the static file handle

    def test_unknown_api_route_returns_json(self):
        response = self.client.get("/api/nope")
        self.assertEqual(response.status_code, 404)
        self.assertIn("error", response.get_json())


class ResumeAnalyzerTests(ApiTestCase):
    def test_missing_file_is_rejected(self):
        response = self.client.post(
            "/api/resume/analyze", data={"target_role": "Backend Developer"}
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("error", response.get_json())

    def test_missing_target_role_is_rejected(self):
        response = self.client.post(
            "/api/resume/analyze",
            data={"resume": (BytesIO(b"%PDF-1.4 fake"), "resume.pdf")},
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("job role", response.get_json()["error"].lower())

    def test_unsupported_extension_is_rejected(self):
        response = self.client.post(
            "/api/resume/analyze",
            data={
                "resume": (BytesIO(b"hello there"), "resume.txt"),
                "target_role": "Backend Developer",
            },
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("pdf", response.get_json()["error"].lower())

    def test_corrupt_pdf_is_reported_friendly(self):
        response = self.client.post(
            "/api/resume/analyze",
            data={
                "resume": (BytesIO(b"not really a pdf at all"), "resume.pdf"),
                "target_role": "Backend Developer",
            },
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 400)
        self.assertTrue(response.get_json()["error"])

    @mock.patch("services.gemini_client.generate_json", return_value=FAKE_ANALYSIS)
    def test_successful_analysis_is_saved(self, _generate):
        response = self.client.post(
            "/api/resume/analyze",
            data={
                "resume": (BytesIO(make_docx_bytes()), "resume.docx"),
                "target_role": "Backend Developer",
            },
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 201)
        record = response.get_json()["analysis"]
        self.assertEqual(record["ats_score"], 72)
        self.assertEqual(record["target_role"], "Backend Developer")
        _generate.assert_called_once()

        listing = self.client.get("/api/resume/analyses").get_json()["analyses"]
        self.assertEqual(len(listing), 1)
        self.assertEqual(listing[0]["id"], record["id"])
        self.assertEqual(listing[0]["score_label"], "Good")

        detail = self.client.get("/api/resume/analyses/" + record["id"])
        self.assertEqual(detail.status_code, 200)
        self.assertEqual(detail.get_json()["analysis"]["result"]["ats_score"], 72)

        deleted = self.client.delete("/api/resume/analyses/" + record["id"])
        self.assertEqual(deleted.status_code, 200)
        self.assertEqual(self.client.get("/api/resume/analyses").get_json()["analyses"], [])

    def test_unknown_analysis_returns_404(self):
        response = self.client.get("/api/resume/analyses/does-not-exist")
        self.assertEqual(response.status_code, 404)
        self.assertIn("error", response.get_json())


class CareerGuideTests(ApiTestCase):
    def test_missing_fields_are_rejected(self):
        response = self.client.post("/api/career/plan", json={"education": "B.Tech"})
        self.assertEqual(response.status_code, 400)
        self.assertIn("error", response.get_json())

    def test_invalid_body_is_rejected(self):
        response = self.client.post(
            "/api/career/plan", data="not json", content_type="application/json"
        )
        self.assertEqual(response.status_code, 400)

    @mock.patch(
        "services.gemini_client.generate_json",
        side_effect=AIRequestError("Gemini is temporarily unavailable.", 503),
    )
    def test_gemini_failure_returns_clean_json(self, _generate):
        response = self.client.post(
            "/api/career/plan",
            json={
                "education": "B.Tech CSE",
                "skills": "Python",
                "target_role": "Backend Developer",
                "experience_level": "Fresher",
                "time_per_day": "1 hour",
            },
        )

        self.assertEqual(response.status_code, 503)
        self.assertEqual(
            response.get_json(), {"error": "Gemini is temporarily unavailable."}
        )

    @mock.patch("services.gemini_client.generate_json", return_value=FAKE_PLAN)
    def test_successful_plan_is_saved(self, _generate):
        response = self.client.post(
            "/api/career/plan",
            json={
                "education": "B.Tech CSE, 3rd year",
                "skills": "Python, HTML, CSS",
                "interests": "web development",
                "target_role": "Backend Developer",
                "experience_level": "Fresher (0-1 years)",
                "time_per_day": "2-3 hours",
            },
        )
        self.assertEqual(response.status_code, 201)
        plan = response.get_json()["plan"]
        self.assertEqual(plan["profile"]["target_role"], "Backend Developer")

        listing = self.client.get("/api/career/plans").get_json()["plans"]
        self.assertEqual(len(listing), 1)
        self.assertEqual(listing[0]["id"], plan["id"])

        deleted = self.client.delete("/api/career/plans/" + plan["id"])
        self.assertEqual(deleted.status_code, 200)
        self.assertEqual(self.client.get("/api/career/plans").get_json()["plans"], [])

    def test_unknown_plan_returns_404(self):
        self.assertEqual(self.client.delete("/api/career/plans/nope").status_code, 404)


class ResumeParserTests(unittest.TestCase):
    def test_docx_text_is_extracted_including_tables(self):
        text = extract_text("resume.docx", make_docx_bytes())
        self.assertIn("Python", text)
        self.assertIn("Anna University", text)

    def test_empty_file_is_rejected(self):
        with self.assertRaises(ResumeParseError):
            extract_text("resume.docx", b"")

    def test_txt_file_is_rejected(self):
        with self.assertRaises(ResumeParseError):
            extract_text("resume.txt", b"some text")

    def test_too_short_docx_is_rejected(self):
        import docx

        document = docx.Document()
        document.add_paragraph("Hi")
        buffer = BytesIO()
        document.save(buffer)

        with self.assertRaises(ResumeParseError):
            extract_text("resume.docx", buffer.getvalue())

    def test_damaged_docx_is_rejected(self):
        with self.assertRaises(ResumeParseError):
            extract_text("resume.docx", b"this is not a zip archive")


if __name__ == "__main__":
    unittest.main(verbosity=2)



